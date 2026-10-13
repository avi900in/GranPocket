import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = docx.Document()

# Page margins
for s in doc.sections:
    s.top_margin = Inches(0.8)
    s.bottom_margin = Inches(0.8)
    s.left_margin = Inches(0.8)
    s.right_margin = Inches(0.8)

# Title
title = doc.add_paragraph()
title_run = title.add_run('GranPocket: Anchor Insurer Strategy & Enterprise Commercialization Blueprint')
title_run.font.size = Pt(20)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(15, 26, 32)

sub = doc.add_paragraph()
sub_run = sub.add_run('Enhanced with Tencent RTC (TRTC) Live Voice Triage & Singapore Integrated Shield Plan Actuarial ROI Model')
sub_run.font.size = Pt(11)
sub_run.font.italic = True
sub_run.font.color.rgb = RGBColor(75, 85, 99)

doc.add_paragraph('—' * 50)

with open('Granpocket Anchor Insurer Strategy.md', 'r') as f:
    text = f.read()

lines = text.split('\n')
in_code = False

for line in lines:
    line_s = line.strip()
    if not line_s:
        continue
    if line_s.startswith('```'):
        in_code = not in_code
        continue
    if in_code:
        p = doc.add_paragraph()
        r = p.add_run(line)
        r.font.size = Pt(8.5)
        r.font.name = 'Courier New'
        p.paragraph_format.space_after = Pt(1)
        continue
    if line_s.startswith('# '):
        continue
    elif line_s.startswith('## '):
        p = doc.add_paragraph()
        r = p.add_run(line_s[3:])
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = RGBColor(26, 54, 93)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
    elif line_s.startswith('### '):
        p = doc.add_paragraph()
        r = p.add_run(line_s[4:])
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(30, 41, 59)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
    elif line_s.startswith('#### '):
        p = doc.add_paragraph()
        r = p.add_run(line_s[5:])
        r.font.size = Pt(10.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(51, 65, 85)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
    elif line_s.startswith('* ') or line_s.startswith('- '):
        p = doc.add_paragraph(style='List Bullet')
        r = p.add_run(line_s[2:].replace('**', ''))
        r.font.size = Pt(9.5)
        p.paragraph_format.space_after = Pt(2)
    elif line_s.startswith('|'):
        p = doc.add_paragraph()
        r = p.add_run(line_s)
        r.font.size = Pt(8.5)
        r.font.name = 'Courier New'
        p.paragraph_format.space_after = Pt(1)
    else:
        p = doc.add_paragraph()
        r = p.add_run(line_s.replace('**', ''))
        r.font.size = Pt(9.5)
        p.paragraph_format.space_after = Pt(3)

doc.save('GranPocket Anchor Insurer Strategy.docx')
print('Successfully regenerated GranPocket Anchor Insurer Strategy.docx')
